"""Сборка презентации PPTX и сопроводительного DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ACCENT = PptRGB(0xF5, 0xA6, 0x23)
DARK = PptRGB(0x0E, 0x11, 0x16)
WHITE = PptRGB(0xF4, 0xF1, 0xEA)
MUTED = PptRGB(0xC5, 0xCB, 0xD3)


def _set_run_font(run, size=18, bold=False, color=WHITE, name="Calibri"):
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _bg(slide, rgb):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def _box(slide, l, t, w, h, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(PInches(l), PInches(t), PInches(w), PInches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=size, bold=bold, color=color)
    return shape


def _bullets(slide, l, t, w, h, items, size=16):
    shape = slide.shapes.add_textbox(PInches(l), PInches(t), PInches(w), PInches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.size = PPt(size)
        p.font.color.rgb = WHITE
        p.font.name = "Calibri"
        p.space_after = PPt(8)
    return shape


def build_pptx():
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)

    def add():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _bg(slide, DARK)
        _box(slide, 0.4, 7.1, 12, 0.3, "StroySync  ·  хакатон  ·  компьютерное зрение на стройплощадке", 11, color=MUTED)
        return slide

    s = add()
    _box(s, 0.7, 1.6, 12, 0.5, "ПРОТОТИП  ·  COMPUTER VISION  ·  CALENDAR MATCHING", 14, True, ACCENT)
    _box(s, 0.7, 2.2, 12, 1.2, "StroySync", 54, True)
    _box(s, 0.7, 3.5, 11, 1.2, "Интеллектуальный контроль строительной площадки:\nтехника на камере ↔ этап календарного графика ↔ объяснимое предупреждение", 22)
    _box(s, 0.7, 5.4, 11, 0.8, "YOLO-World  ·  FastAPI  ·  React  ·  PostgreSQL  ·  Docker", 16, color=ACCENT)

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Задача", 32, True, ACCENT)
    _bullets(s, 0.7, 1.3, 12, 5.5, [
        "По снимкам камер понять, какая техника реально работает на площадке.",
        "Связать кадр с иерархическим календарным графиком (WBS, зоны, сроки).",
        "Найти отклонения: нет обязательных машин, чужой этап, работы вперёд/назад графика.",
        "Каждое предупреждение должно быть понятным прорабу: что видно, какой этап, какое правило.",
        "Показать, как ставить камеры, чтобы аналитика была достоверной.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Концепция", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Камера — это не «просто глаз», а датчик конкретной зоны графика.",
        "Этап работ описывается цифровым отпечатком техники: обязательно / можно / нельзя.",
        "Снимок превращается во множество классов. Jaccard к отпечатку = степень соответствия.",
        "Сравнение с прошлым и будущим WBS выявляет отставание и опережение.",
        "Правила именованы (MISSING_REQUIRED, FORBIDDEN_EQUIPMENT, AHEAD_OF_SCHEDULE…) — вывод проверяем.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Архитектура", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Клиент: React-дашборд (мониторинг, график, камеры, история).",
        "API: FastAPI — загрузка снимка, детекция, сопоставление, CRUD результатов.",
        "CV: YOLO-World по текстовым промптам 8 типов техники; запасной YOLOv8n; demo для питча.",
        "Matching engine: база знаний этапов + зона камеры + дата снимка.",
        "Хранение: PostgreSQL/SQLite — техника, WBS, камеры, снимки, детекции, отклонения.",
        "Поставка: Docker Compose (frontend + backend + Postgres).",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Почему такой стек", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "YOLO-World: 8 классов ТЗ без ожидания размеченного датасета организаторов; легко добавить типы.",
        "Дообучение классического YOLO — следующий шаг, когда появится разметка площадки.",
        "FastAPI: типизированный API, файлы, OpenAPI для жюри (/docs).",
        "React: операторский интерфейс, а не только блокнот Streamlit.",
        "PostgreSQL: реляционный график, камеры, история анализов; SQLite для ноутбука.",
        "Docker: воспроизводимый запуск на машине жюри.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Конвейер обнаружения", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Вход: кадр + camera_id + timestamp (ключ к графику).",
        "Промпты: dump truck, excavator, road roller, knuckle boom crane, concrete mixer, bulldozer, truck, mobile crane.",
        "NMS по классу, порог уверенности, подписи на русском на аннотированном кадре.",
        "Ограничение: модель открытая, не fine-tune на площадке заказчика. Demo-режим гарантирует питч.",
        "Расширение: дообучение на архиве организаторов + tracking (ByteTrack) на видеопотоке.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Снимок ↔ этап графика", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "CAM-01 смотрит в котлован → WBS 2.*, 3.* (земляные и фундамент).",
        "CAM-03 смотрит на пятно застройки → WBS 4.*, 5.* (каркас, кровля).",
        "На 17.09.2026 активны фундамент (3 / 3.1) и старт каркаса (4 / 4.1).",
        "Если CAM-01 видит каток — срабатывает FORBIDDEN_EQUIPMENT: каток принадлежит этапу благоустройства.",
        "Если состав ближе к будущему WBS, чем к текущему — AHEAD_OF_SCHEDULE.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Предупреждения", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "MISSING_REQUIRED — этап идёт, обязательной техники нет (простой / отставание / слепая зона).",
        "FORBIDDEN_EQUIPMENT — машина с чужого этапа, критический сигнал.",
        "UNEXPECTED_EQUIPMENT — не запрещено, но нехарактерно (транзит по кадру).",
        "AHEAD / BEHIND — отпечаток техники похож на будущий или прошедший WBS.",
        "NO_ACTIVITY — пустой кадр при активном этапе.",
        "ON_TRACK — обязательный набор закрыт, запрещённых нет. Это тоже результат анализа.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Демонстрация прототипа", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Веб-интерфейс: http://localhost:5173 (или :8080 в Docker).",
        "Сцены: фундамент в графике, земляные (архив), каркас, смесь этапов, каток-нарушитель, пустой кадр.",
        "Справа — активный WBS, score, текст предупреждения со ссылкой на правило.",
        "Вкладки «График» и «Камеры» показывают ту же базу знаний, что использует matcher.",
        "API и CLI позволяют жюри прогнать свой снимок без UI.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Как ставить камеры", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Одна камера = одна зона графика. Иначе сопоставление неоднозначно.",
        "Высота 8–18 м, угол 30–50°, техника ≥ 8–12% ширины кадра.",
        "Обязательный ракурс на въезде: тип машины виден до слепой зоны котлована.",
        "Перекрытие соседних камер 15–20%. ИК/подсветка для ночных заливок.",
        "Вдоль оси дороги — иначе каток виден торцом и путается с грузовиком.",
        "Не смотреть в низкое зимнее солнце; козырёк и обогрев зимой.",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Масштабирование", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "RTSP + ByteTrack: техника не «мелькает», а живёт как объект во времени.",
        "Fine-tune YOLO на датасете площадки (сезон, ночь, грязь, частичные перекрытия).",
        "Импорт графика из MS Project / Primavera / 1С, несколько объектов холдинга.",
        "Edge (Jetson) на бытовке: в облако уходят события, не видео.",
        "Цифровой двойник: статус WBS красится по факту с камер, не по отчёту подрядчика.",
        "ANPR на въезде + наряд-допуск: «миксер не заявлен на сегодня».",
    ])

    s = add()
    _box(s, 0.7, 0.4, 12, 0.6, "Ограничения прототипа", 32, True, ACCENT)
    _bullets(s, 0.7, 1.2, 12, 5.6, [
        "Open-vocabulary модель не заменяет разметку конкретной площадки.",
        "Синтетические демо-кадры иллюстрируют логику; боевое качество — на реальном архиве.",
        "Нет трекинга и работы по видеопотоку в этой версии.",
        "График — учебный жилой объект, июнь 2026 — март 2027.",
        "Юридически значимый контроль требует регламента заказчика и человека в контуре.",
    ])

    s = add()
    _box(s, 0.7, 2.0, 12, 0.5, "Итог", 18, True, ACCENT)
    _box(s, 0.7, 2.6, 12, 1.4, "StroySync делает график проверяемым\nпо тому, что реально стоит в кадре.", 28, True)
    _box(s, 0.7, 4.5, 12, 1.2, "Детекция → зона камеры → активный WBS → правило → объяснение.", 18, color=MUTED)

    path = DOCS / "StroySync_Presentation.pptx"
    prs.save(path)
    print("wrote", path)


def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    def h(text, level=1):
        doc.add_heading(text, level=level)

    def p(text):
        doc.add_paragraph(text)

    h("StroySync — сопроводительная документация")
    p("Прототип системы компьютерного зрения для сопоставления снимков строительных площадок с календарным графиком работ. Документ описывает модель данных, ограничения моделей, логику сопоставления и инструкции по запуску.")

    h("1. Назначение")
    p("Система принимает растровый снимок с камеры видеонаблюдения, обнаруживает строительную технику, определяет, какие этапы WBS должны быть активны на дату кадра в зоне этой камеры, и формирует статус площадки плюс список объяснимых отклонений.")

    h("2. Модель данных")
    h("2.1. Типы техники (equipment_types)", 2)
    p("Первичный ключ — code. Хранятся русское и английское имена, цветовая метка для отрисовки и список текстовых промптов для YOLO-World.")
    p("Коды: dump_truck (самосвал), excavator (экскаватор), roller (каток), manipulator_crane (кран-манипулятор), concrete_mixer (бетоносмеситель), bulldozer (бульдозер), truck (грузовик), mobile_crane (автокран). Перечень совпадает с ТЗ и расширяется добавлением записи в каталог и промпта, без переписывания детектора.")

    h("2.2. Этапы работ (work_stages)", 2)
    p("Иерархический WBS: wbs, parent_wbs, name, level, zone, start_date, end_date. Для сопоставления у этапа заданы три множества техники: required_equipment, allowed_equipment, forbidden_equipment и typical_signature — «цифровой отпечаток» типичного состава. Описание (description) используется в UI и в текстах предупреждений.")

    h("2.3. Камеры (cameras)", 2)
    p("id, name, zone, linked_wbs (список кодов этапов, которые камера должна видеть), view, height_m, angle_deg, recommendations. Связь снимка с графиком всегда идёт через камеру: без camera_id этап не выбирается.")

    h("2.4. Снимки, детекции, отклонения", 2)
    p("snapshots: исходный и аннотированный файлы, timestamp, detector_backend, site_status (on_track / warning / critical / unknown), match_score, summary, JSON active_stages и detected_counts.")
    p("detections: класс, уверенность, bbox (x1,y1,x2,y2) в пикселях исходного кадра.")
    p("deviations: code правила, severity, title, explanation, wbs, evidence (JSON со списками missing/forbidden/scores). Это и есть журнал нарушений/отклонений.")

    h("3. Код и модули")
    p("backend/app/services/detector.py — загрузка YOLO-World / YOLOv8 / demo.")
    p("backend/app/services/knowledge.py — таксономия и Jaccard.")
    p("backend/app/services/matcher.py — правила сопоставления и тексты предупреждений. Нетривиальная логика прокомментирована в модуле.")
    p("backend/app/services/pipeline.py — транзакция: детект → матчинг → аннотация → запись в БД.")
    p("backend/app/seed.py — эталонный график и камеры.")
    p("frontend/src/App.jsx — операторский дашборд.")

    h("4. Условия и ограничения модели")
    p("Детектор в боевом режиме — открытая модель YOLO-World (yolov8s-worldv2) компании Ultralytics. Она не обучалась на архиве организаторов хакатона. Классы задаются текстовыми промптами на английском; русские синонимы используются при разборе подписи.")
    p("Запасной режим YOLOv8n обучен на COCO и уверенно видит только обобщённый класс truck/bus. Экскаватор, каток, автокран, манипулятор через COCO не выделяются — поэтому COCO не основной путь.")
    p("Demo-режим не является нейросетью: по ключу в имени файла (foundation, earthworks, frame, mismatch, roads, empty) возвращаются стабильные боксы. Он нужен для питча, автотестов и машин без GPU/весов.")
    p("Распознаваемые типы — восемь позиций ТЗ. Сторонние датасеты (SODA, ACID, Roboflow construction) допускаются для дообучения; в прототипе веса fine-tune не прилагаются, чтобы репозиторий оставался лёгким.")
    p("Логика сопоставления настроена на учебный календарь жилого объекта: подготовка → земляные → фундамент → каркас → кровля/фасад → благоустройство. Дата «сегодня» в демо — 17 сентября 2026, активны фундамент и старт каркаса. Для другого графика достаточно загрузить новые work_stages.")

    h("5. Алгоритм сопоставления")
    p("1) По captured_at выбираются этапы, у которых start_date ≤ дата ≤ end_date.")
    p("2) Оставляются этапы, чьи wbs связаны с камерой (точное совпадение или префикс).")
    p("3) Для активного этапа: missing = required − detected; unexpected = detected − allowed; forbidden = detected ∩ forbidden.")
    p("4) Jaccard(detected, typical_signature) считается для активного, прошлого и будущего этапа. Если будущий отпечаток заметно ближе — AHEAD_OF_SCHEDULE, если прошлый — BEHIND_SCHEDULE.")
    p("5) Статус площадки: critical при запрещённой технике; warning при пропусках/отставании/пустом кадре; on_track если обязательный набор закрыт.")
    p("Каждое предупреждение содержит WBS, код правила и evidence — его можно проверить глазами по аннотированному снимку и таблице графика.")

    h("6. Сборка, установка и запуск")
    h("6.1. Локально, Windows", 2)
    p("1. Установите Python 3.11+ и Node.js 20+.")
    p("2. python -m venv .venv и .\\.venv\\Scripts\\Activate.ps1")
    p("3. pip install -r backend\\requirements.txt")
    p("4. python scripts\\generate_samples.py — кадры и Excel-график.")
    p("5. $env:PYTHONPATH = \"<корень>\\backend\"; $env:DETECTOR_BACKEND = \"demo\"")
    p("6. uvicorn app.main:app --app-dir backend --reload --port 8000")
    p("7. В другом терминале: cd frontend; npm install; npm run dev")
    p("8. Браузер: http://localhost:5173  API: http://localhost:8000/docs")

    h("6.2. Docker", 2)
    p("docker compose up --build")
    p("UI: http://localhost:8080  API: http://localhost:8000")
    p("По умолчанию DETECTOR_BACKEND=demo. Для YOLO-World смените переменную у сервиса backend (скачаются веса PyTorch, образ тяжелее, первый запуск дольше).")

    h("6.3. Проверка чужого снимка", 2)
    p("POST /api/analyze (multipart: file, camera_id, captured_at). Либо CLI: python -m app.cli analyze photo.jpg --camera CAM-01 --at 2026-09-17T10:00:00")

    h("6.4. Тесты", 2)
    p("cd backend && pytest -q — проверяют matcher без нейросети: «в графике», запрещённый каток, сигнал опережения.")

    h("7. Рекомендации по камерам")
    p("Камера должна быть датчиком зоны WBS, а не обзорной «на всякий случай». Высота 8–18 м, угол 30–50°, объект техники 8–12% ширины кадра, отдельный ракурс на въезде, перекрытие 15–20%, ИК для ночных бетонирований, ориентация вдоль оси для катков, защита от низкого солнца и осадков. Подробности вынесены в API /api/camera-guide и на вкладку «Камеры».")

    h("8. Развитие")
    p("Дообучение на архиве организаторов; видеопоток RTSP + ByteTrack; импорт графика из MS Project/Primavera; мультиобъект; edge-инференс; наряд-допуски и ANPR; окраска цифрового двойника по факту с камер.")

    path = DOCS / "StroySync_Documentation.docx"
    doc.save(path)
    print("wrote", path)


if __name__ == "__main__":
    DOCS.mkdir(parents=True, exist_ok=True)
    build_pptx()
    build_docx()
